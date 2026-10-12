"""
utils/file_io.py
================

Read uploaded files into memory.

Two kinds of input are supported:

* Images: decoded with OpenCV into NumPy arrays.
* Text documents: decoded as UTF-8 or, for .docx, read with python-docx.

Centralizing this means every section applies the same checks and the
same error messages.
"""

from pathlib import Path

import cv2
import numpy as np

# Project root folder. Used by the File section to locate sample data.
PROJECT_ROOT = Path(__file__).resolve().parents[1]


def decode_image(raw_bytes):
    """
    Decode uploaded bytes into a NumPy image.

    Alpha is preserved if the file has it; grayscale images are upgraded
    to three channels for uniform downstream handling.
    """
    image = cv2.imdecode(np.frombuffer(raw_bytes, np.uint8),
                         cv2.IMREAD_UNCHANGED)
    if image is None:
        raise ValueError("This file could not be decoded as an image.")
    if image.dtype != np.uint8:
        raise ValueError("Only 8-bit images are supported.")
    if image.shape[0] * image.shape[1] > 25_000_000:
        raise ValueError("Image is too large; keep it under 25 megapixels.")
    if image.ndim == 2:
        image = cv2.cvtColor(image, cv2.COLOR_GRAY2BGR)
    return image


def encode_image(image, extension):
    """Encode a NumPy image into bytes for download."""
    if image.dtype != np.uint8:
        image = np.clip(image * 255, 0, 255).astype(np.uint8)
    ok, buffer = cv2.imencode("." + extension, image)
    if not ok:
        raise ValueError("Image encoding failed.")
    return buffer.tobytes()


def read_text(uploaded_file):
    """Read a text file (UTF-8) or a .docx file into a single string."""
    if uploaded_file.name.lower().endswith(".docx"):
        from docx import Document
        uploaded_file.seek(0)
        document = Document(uploaded_file)
        paragraphs = [p.text for p in document.paragraphs]
        table_lines = [
            " | ".join(cell.text for cell in row.cells)
            for table in document.tables
            for row in table.rows
        ]
        return "\n".join(paragraphs + table_lines)
    return uploaded_file.getvalue().decode("utf-8-sig")
